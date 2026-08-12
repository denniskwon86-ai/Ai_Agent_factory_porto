"""기술·제품 용어 인벤토리를 전환 사전(JSON/XLSX)으로 빌드한다.

정본 입력은 사람이 검토하는 Markdown 인벤토리다. Excel과 프론트엔드가 각자 용어를
복사해 갖지 않도록 이 스크립트가 두 산출물을 함께 만든다.
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "architecture" / "TECHNOLOGY_TERMINOLOGY_REDEFINITION_INVENTORY_2026-08-11.md"
POLICY_SOURCE = ROOT / "data" / "terminology" / "technology_terminology_policy.json"
CANONICAL_JSON = ROOT / "data" / "terminology" / "technology_terminology_glossary.json"
FRONTEND_JSON = ROOT / "frontend" / "src" / "data" / "technologyTerminologyGlossary.json"
XLSX = ROOT / "docs" / "architecture" / "AI_FACTORY_STUDIO_TECHNOLOGY_TERMINOLOGY_DICTIONARY_2026-08-12.xlsx"

AREA_BY_SECTION = {
    2: "제품 정체성·최상위 구조",
    3: "회사·조직·문맥·권한 범위",
    4: "데이터 기반·MDM·카탈로그",
    5: "지식·근거·검색",
    6: "AI 에이전트·모델·워크플로우",
    7: "업무 앱 생성·오케스트레이션·산출물",
    8: "업무 앱 런타임·보안·실시간 통신",
    9: "경영계획·시뮬레이션·디지털트윈",
    10: "의사결정·협업·실행·발간",
    11: "전역 AI 비서·감독 기능",
    12: "외부 연계·MCP·수집",
    13: "품질·감사·운영·테스트",
    14: "주요 화면·내비게이션",
}

# 사용자에게 먼저 보여 줄 이름과 기술 설명에서 보존할 이름을 분리한다. 목록에 없는 항목은
# 원문을 보존하면서 상태를 명시하므로, 아직 합의되지 않은 새 이름이 확정된 것처럼 퍼지지 않는다.
OVERRIDES: dict[str, dict[str, str]] = {
    "PROD-01": {"user": "AI Factory Studio", "english": "AI Factory Studio", "tech": "ai_factory_studio", "status": "현행 유지"},
    "PROD-02": {"user": "제조 경영 시스템", "english": "Manufacturing Management System", "tech": "management_system", "status": "권장안"},
    "PROD-03": {"user": "제조 경영 디지털트윈", "english": "Manufacturing Management Twin", "tech": "management_twin", "status": "권장안"},
    "PROD-04": {"user": "제조 경영 디지털트윈", "english": "Manufacturing Management Twin", "tech": "management_twin", "status": "권장안"},
    "PROD-05": {"user": "전사 경영 허브", "english": "Enterprise Management Hub", "tech": "enterprise_home", "status": "과거 별칭"},
    "PROD-06": {"user": "전사 경영 허브", "english": "Enterprise Management Hub", "tech": "enterprise_home", "status": "권장안"},
    "PROD-07": {"user": "경영진 관제 화면", "english": "Executive Cockpit", "tech": "executive_cockpit", "status": "권장안"},
    "PROD-08": {"user": "업무 앱 스튜디오", "english": "Operational App Studio", "tech": "app_studio", "status": "과거 별칭"},
    "PROD-09": {"user": "업무 앱 스튜디오", "english": "Operational App Studio", "tech": "app_factory", "status": "권장안"},
    "PROD-10": {"user": "업무 앱 스튜디오", "english": "Host-Governed Operational App Factory", "tech": "host_governed_app_factory", "status": "기술 전용"},
    "PROD-11": {"user": "업무 앱 만들기", "english": "Build an Operational App", "tech": "app_factory", "status": "권장안"},
    "PROD-12": {"user": "업무 앱 스튜디오", "english": "Operational App Factory", "tech": "app_factory", "status": "권장안"},
    "PROD-13": {"user": "플랫폼 내 업무 앱", "english": "App-in-App Runtime", "tech": "app_in_app", "status": "권장안"},
    "PROD-14": {"user": "생성된 업무 앱", "english": "Generated Operational App", "tech": "generated_app", "status": "권장안"},
    "PROD-15": {"user": "업무 앱", "english": "Operational App", "tech": "operational_app", "status": "권장안"},
    "PROD-16": {"user": "업무 프로그램", "english": "Business Program", "tech": "business_program", "status": "권장안"},
    "PROD-17": {"user": "앱 제작 프로젝트", "english": "App Build Project", "tech": "project", "status": "권장안"},
    "PROD-18": {"user": "통합 제작 프로젝트", "english": "Multi-App Program", "tech": "mega_project", "status": "권장안"},
    "PROD-19": {"user": "부서 업무공간", "english": "Department Workspace", "tech": "workspace", "status": "권장안"},
    "ORG-06": {"user": "조직 단위", "english": "Organization Node", "tech": "org_node", "status": "기술 전용"},
    "ORG-07": {"user": "현재 회사 문맥", "english": "Enterprise Context", "tech": "enterprise_context", "status": "권장안"},
    "ORG-08": {"user": "현재 회사 문맥", "english": "Company Context", "tech": "company_context", "status": "권장안"},
    "ORG-09": {"user": "현재 작업 문맥", "english": "Active Context", "tech": "active_context", "status": "기술 전용"},
    "ORG-10": {"user": "회사·조직 선택", "english": "Organization Context Switcher", "tech": "context_switcher", "status": "권장안"},
    "ORG-11": {"user": "고객 격리 단위", "english": "Tenant", "tech": "tenant_id", "status": "기술 전용"},
    "ORG-12": {"user": "접근 범위", "english": "Authorization Scope", "tech": "scope", "status": "검토 필요"},
    "ORG-13": {"user": "조직 접근 범위", "english": "Enterprise Scope", "tech": "enterprise_scope_id", "status": "권장안"},
    "ORG-20": {"user": "가상 회사", "english": "Virtual Company", "tech": "virtual_company", "status": "권장안"},
    "ORG-24": {"user": "회사 브랜드 설정", "english": "Company Brand Profile", "tech": "company_brand_profile", "status": "권장안"},
    "ORG-25": {"user": "회사 복제", "english": "Company Clone", "tech": "company_clone", "status": "권장안"},
    "DATA-01": {"user": "기준정보 관리", "english": "Master Data Management", "tech": "mdm", "status": "권장안"},
    "DATA-02": {"user": "기준정보", "english": "Master Data", "tech": "master_data", "status": "권장안"},
    "DATA-03": {"user": "기준정보", "english": "Master Data", "tech": "master_data", "status": "권장안"},
    "DATA-04": {"user": "코드·참조정보", "english": "Reference Data", "tech": "reference_data", "status": "권장안"},
    "DATA-05": {"user": "공식 근거 등록부", "english": "Reference Registry", "tech": "reference_registry", "status": "권장안"},
    "DATA-06": {"user": "데이터 원천 등록부", "english": "Source Registry", "tech": "source_registry", "status": "권장안"},
    "DATA-07": {"user": "데이터 카탈로그", "english": "Data Catalog", "tech": "data_catalog", "status": "현행 유지"},
    "DATA-13": {"user": "데이터 항목 정의서", "english": "Data Dictionary", "tech": "data_dictionary", "status": "권장안"},
    "DATA-14": {"user": "업무 용어 사전", "english": "Business Glossary", "tech": "business_glossary", "status": "권장안"},
    "DATA-16": {"user": "코드·항목 연결표", "english": "Crosswalk", "tech": "crosswalk", "status": "권장안"},
    "DATA-18": {"user": "업무표준", "english": "Work Standard", "tech": "work_standard", "status": "현행 유지"},
    "DATA-19": {"user": "데이터 계보", "english": "Data Lineage", "tech": "lineage", "status": "권장안"},
    "DATA-20": {"user": "데이터 품질", "english": "Data Quality", "tech": "data_quality", "status": "현행 유지"},
    "DATA-22": {"user": "수집 원본", "english": "Raw", "tech": "RAW", "status": "권장안"},
    "DATA-23": {"user": "검증 완료", "english": "Validated", "tech": "VALIDATED", "status": "권장안"},
    "DATA-24": {"user": "의사결정 사용 승인", "english": "Certified", "tech": "CERTIFIED", "status": "권장안"},
    "DATA-25": {"user": "데모 사용 승인", "english": "Certified for Demo", "tech": "CERTIFIED_FOR_DEMO", "status": "권장안"},
    "DATA-26": {"user": "검증 격리", "english": "Quarantine", "tech": "QUARANTINE", "status": "권장안"},
    "DATA-28": {"user": "실적", "english": "Actual", "tech": "ACTUAL", "status": "권장안"},
    "DATA-29": {"user": "계획", "english": "Plan", "tech": "PLAN", "status": "권장안"},
    "DATA-30": {"user": "전망", "english": "Forecast", "tech": "FORECAST", "status": "권장안"},
    "DATA-31": {"user": "시나리오", "english": "Scenario", "tech": "SCENARIO", "status": "권장안"},
    "DATA-32": {"user": "합성 데이터", "english": "Synthetic Data", "tech": "SYNTHETIC", "status": "권장안"},
    "DATA-41": {"user": "업무 시작 키트", "english": "Business Starter Kit", "tech": "starter_kit", "status": "권장안"},
    "KNOW-01": {"user": "전사 지식 허브", "english": "Enterprise Knowledge Hub", "tech": "knowledge_hub", "status": "권장안"},
    "KNOW-02": {"user": "지식 저장소", "english": "Knowledge Base", "tech": "knowledge_base", "status": "기술 전용"},
    "KNOW-03": {"user": "지식팩", "english": "Knowledge Pack", "tech": "knowledge_pack", "status": "권장안"},
    "KNOW-05": {"user": "근거 자료", "english": "Evidence", "tech": "evidence", "status": "권장안"},
    "KNOW-07": {"user": "근거 연결", "english": "Grounding", "tech": "grounding", "status": "기술 전용"},
    "KNOW-08": {"user": "근거 기반 생성", "english": "Retrieval-Augmented Generation", "tech": "rag", "status": "기술 전용"},
    "KNOW-09": {"user": "관계 기반 지식 검색", "english": "Graph RAG", "tech": "graph_rag", "status": "검토 필요"},
    "KNOW-10": {"user": "자연어 지식 검색", "english": "Semantic Search", "tech": "semantic_search", "status": "권장안"},
    "AI-01": {"user": "AI 에이전트", "english": "AI Agent", "tech": "agent", "status": "권장안"},
    "AI-02": {"user": "에이전트 등록부", "english": "Agent Registry", "tech": "agent_registry", "status": "권장안"},
    "AI-03": {"user": "에이전트 자산", "english": "Agent Asset", "tech": "agent_asset", "status": "권장안"},
    "AI-04": {"user": "에이전트 팩", "english": "Agent Pack", "tech": "agent_pack", "status": "권장안"},
    "AI-06": {"user": "에이전트 통제소", "english": "Agent Control Center", "tech": "agent_master", "status": "권장안"},
    "AI-07": {"user": "에이전트 자산 통제", "english": "Agent Governance", "tech": "agent_governance", "status": "권장안"},
    "AI-09": {"user": "에이전트 스킬", "english": "Agent Skill", "tech": "skill", "status": "권장안"},
    "AI-10": {"user": "AI 스킬 개선", "english": "Skill Evolution", "tech": "skill_evolution", "status": "권장안"},
    "AI-13": {"user": "AI 작업 흐름", "english": "Agent Workflow", "tech": "workflow", "status": "권장안"},
    "AI-14": {"user": "워크플로우 템플릿", "english": "Workflow Template", "tech": "workflow_template", "status": "권장안"},
    "AI-15": {"user": "실행 그래프", "english": "Agent Graph", "tech": "agent_graph", "status": "기술 전용"},
    "AI-21": {"user": "LLM 통합 게이트웨이", "english": "LLM Gateway", "tech": "llm_gateway", "status": "기술 전용"},
    "AI-24": {"user": "모델 선택 정책", "english": "Model Routing", "tech": "model_routing", "status": "기술 전용"},
    "AI-25": {"user": "모델 대체 경로", "english": "Fallback", "tech": "fallback", "status": "기술 전용"},
    "AI-28": {"user": "사용 모델·비용 현황", "english": "LLM Telemetry", "tech": "llm_telemetry", "status": "권장안"},
    "AI-30": {"user": "비용·품질 최적화", "english": "Cost/Quality Optimizer", "tech": "cost_quality_optimizer", "status": "권장안"},
    "FLOW-01": {"user": "RFP", "english": "Request for Proposal", "tech": "rfp", "status": "표준 약어 유지"},
    "FLOW-02": {"user": "요구사항 구체화", "english": "Clarification", "tech": "clarification", "status": "권장안"},
    "FLOW-03": {"user": "PRD", "english": "Product Requirements Document", "tech": "prd", "status": "표준 약어 유지"},
    "FLOW-06": {"user": "화면 시각검증", "english": "Vision QA", "tech": "vision_qa", "status": "기술 전용"},
    "FLOW-07": {"user": "WBS", "english": "Work Breakdown Structure", "tech": "wbs", "status": "표준 약어 유지"},
    "FLOW-08": {"user": "작업 항목", "english": "Task", "tech": "task", "status": "권장안"},
    "FLOW-09": {"user": "작업 단계", "english": "Stage", "tech": "stage", "status": "권장안"},
    "FLOW-11": {"user": "실행 단위", "english": "Sprint", "tech": "sprint", "status": "기술 전용"},
    "FLOW-14": {"user": "워크플로우 실행기", "english": "Orchestrator", "tech": "orchestrator", "status": "기술 전용"},
    "FLOW-17": {"user": "사용자 검토 관문", "english": "Human-on-the-loop", "tech": "hotl", "status": "기술 전용"},
    "FLOW-18": {"user": "사용자 검토 관문", "english": "User Review Gate", "tech": "user_review_gate", "status": "권장안"},
    "FLOW-19": {"user": "승인 관문", "english": "Approval Gate", "tech": "approval_gate", "status": "권장안"},
    "FLOW-25": {"user": "산출물", "english": "Artifact", "tech": "artifact", "status": "권장안"},
    "FLOW-28": {"user": "보고서 산출물", "english": "Report Output", "tech": "report_output", "status": "권장안"},
    "FLOW-35": {"user": "자가복구", "english": "Self-Healing", "tech": "self_healing", "status": "권장안"},
    "FLOW-37": {"user": "릴리스", "english": "Release", "tech": "release", "status": "권장안"},
    "FLOW-38": {"user": "운영 자산 승격", "english": "Promotion", "tech": "promotion", "status": "권장안"},
    "RUN-01": {"user": "업무 앱 실행 기반", "english": "Host Runtime", "tech": "host_runtime", "status": "기술 전용"},
    "RUN-02": {"user": "플랫폼 연계 SDK", "english": "Host SDK", "tech": "host_sdk", "status": "기술 전용"},
    "RUN-03": {"user": "앱 데이터 통제 경로", "english": "App Data Plane", "tech": "app_data_plane", "status": "기술 전용"},
    "RUN-04": {"user": "앱 권한 선언서", "english": "Capability Manifest", "tech": "capability_manifest", "status": "권장안"},
    "RUN-05": {"user": "앱 선언서", "english": "App Manifest", "tech": "app_manifest", "status": "권장안"},
    "RUN-23": {"user": "실시간 알림", "english": "Server-Sent Events", "tech": "sse", "status": "기술 전용"},
    "RUN-24": {"user": "실시간 연결 인증권", "english": "SSE Ticket", "tech": "sse_ticket", "status": "기술 전용"},
    "RUN-28": {"user": "앱 전달", "english": "App Delivery", "tech": "app_delivery", "status": "권장안"},
    "RUN-29": {"user": "내 앱", "english": "My Apps", "tech": "app_pocket", "status": "권장안"},
    "TWIN-01": {"user": "디지털트윈", "english": "Digital Twin", "tech": "digital_twin", "status": "검토 필요"},
    "TWIN-02": {"user": "제조 경영 디지털트윈", "english": "Manufacturing Management Twin", "tech": "management_twin", "status": "권장안"},
    "TWIN-03": {"user": "경영 시뮬레이터", "english": "Management Simulator", "tech": "simulator", "status": "권장안"},
    "TWIN-04": {"user": "경영 시뮬레이션", "english": "Management Simulation", "tech": "simulation", "status": "권장안"},
    "TWIN-05": {"user": "경영계획 계산 엔진", "english": "Planning Engine", "tech": "planning_engine", "status": "권장안"},
    "TWIN-10": {"user": "조정 변수", "english": "Lever", "tech": "lever", "status": "권장안"},
    "TWIN-11": {"user": "외부환경 지표", "english": "External Indicator", "tech": "external_indicator", "status": "권장안"},
    "TWIN-13": {"user": "비교 기준", "english": "Benchmark", "tech": "benchmark", "status": "권장안"},
    "TWIN-15": {"user": "계산 모델", "english": "Calculation Graph", "tech": "calculation_graph", "status": "권장안"},
    "TWIN-17": {"user": "결정론적 계산 엔진", "english": "Deterministic Engine", "tech": "deterministic_engine", "status": "권장안"},
    "TWIN-29": {"user": "민감도 분석", "english": "Sensitivity Analysis", "tech": "sensitivity_analysis", "status": "권장안"},
    "DEC-01": {"user": "의사결정 안건", "english": "Decision Case", "tech": "decision_case", "status": "권장안"},
    "DEC-02": {"user": "의사결정 검토서", "english": "Decision Package", "tech": "decision_package", "status": "권장안"},
    "DEC-03": {"user": "의사결정 원장", "english": "Decision Ledger", "tech": "decision_ledger", "status": "권장안"},
    "DEC-04": {"user": "의사결정 센터", "english": "Decision Center", "tech": "decision_center", "status": "권장안"},
    "DEC-05": {"user": "의사결정 요청", "english": "Decision Request", "tech": "decision_request", "status": "권장안"},
    "DEC-13": {"user": "효과 측정", "english": "Effect Measurement", "tech": "effect_measurement", "status": "권장안"},
    "DEC-14": {"user": "폐루프 운영", "english": "Closed Loop", "tech": "closed_loop", "status": "권장안"},
    "DEC-16": {"user": "보고서 발간", "english": "Publication", "tech": "publication", "status": "권장안"},
    "DEC-20": {"user": "경영진 브리핑", "english": "Executive Briefing", "tech": "executive_briefing", "status": "권장안"},
    "DEC-23": {"user": "앱 전달", "english": "App Delivery", "tech": "app_push", "status": "권장안"},
    "ASST-01": {"user": "Jarvis", "english": "Jarvis", "tech": "product_assistant", "status": "권장안"},
    "ASST-02": {"user": "Jarvis", "english": "Jarvis", "tech": "product_assistant", "status": "과거 별칭"},
    "ASST-03": {"user": "품질 감독 에이전트", "english": "Pipeline Quality Supervisor", "tech": "supervisor_agent", "status": "권장안"},
    "ASST-04": {"user": "전역 AI 비서", "english": "Enterprise AI Assistant", "tech": "global_supervisor", "status": "검토 필요"},
    "ASST-05": {"user": "백그라운드 감시기", "english": "Supervisor Daemon", "tech": "supervisor_daemon", "status": "기술 전용"},
    "ASST-07": {"user": "업무·데이터 설계 상담", "english": "Business and Data Design Advisor", "tech": "advisor", "status": "권장안"},
    "INT-01": {"user": "외부 시스템 연결", "english": "Connector", "tech": "connector", "status": "권장안"},
    "INT-02": {"user": "연결 등록부", "english": "Connector Registry", "tech": "connector_registry", "status": "기술 전용"},
    "INT-04": {"user": "MCP 연계", "english": "Model Context Protocol", "tech": "mcp", "status": "기술 전용"},
    "INT-07": {"user": "기존 업무 시스템", "english": "Legacy System", "tech": "legacy_system", "status": "권장안"},
    "INT-08": {"user": "외부 참여자 업무 시스템", "english": "External Engagement System", "tech": "external_engagement_system", "status": "권장안"},
    "INT-10": {"user": "LPL(회사별 사례)", "english": "LPL", "tech": "lpl", "status": "과거 별칭"},
    "OPS-01": {"user": "거버넌스", "english": "Governance", "tech": "governance", "status": "검토 필요"},
    "OPS-02": {"user": "거버넌스 현황", "english": "Governance Console", "tech": "governance_console", "status": "권장안"},
    "OPS-03": {"user": "시스템 관리", "english": "System Administration", "tech": "admin_console", "status": "권장안"},
    "OPS-08": {"user": "품질 결과", "english": "Quality Telemetry", "tech": "quality_telemetry", "status": "권장안"},
    "OPS-15": {"user": "릴리스 적격성 관문", "english": "Release Gate", "tech": "release_gate", "status": "권장안"},
    "OPS-16": {"user": "구문 검사", "english": "Syntax Check", "tech": "syntax_check", "status": "기술 전용"},
    "OPS-20": {"user": "백엔드 기동 검사", "english": "Backend Smoke Test", "tech": "backend_smoke_test", "status": "기술 전용"},
    "OPS-25": {"user": "추적성", "english": "Traceability", "tech": "traceability", "status": "권장안"},
    "OPS-26": {"user": "추적성 그래프", "english": "Traceability Graph", "tech": "traceability_graph", "status": "권장안"},
    "UI-01": {"user": "전사 경영 허브", "english": "Enterprise Management Hub", "tech": "enterprise_page", "status": "권장안"},
    "UI-02": {"user": "업무 앱 만들기", "english": "Build an Operational App", "tech": "build_sw", "status": "권장안"},
    "UI-03": {"user": "업무 앱", "english": "Operational Apps", "tech": "operate", "status": "권장안"},
    "UI-04": {"user": "경영 시뮬레이션", "english": "Management Simulation", "tech": "simulate", "status": "권장안"},
    "UI-05": {"user": "전사 지식 허브", "english": "Enterprise Knowledge Hub", "tech": "knowledge", "status": "권장안"},
    "UI-06": {"user": "데이터 기반", "english": "Enterprise Data Foundation", "tech": "data", "status": "권장안"},
    "UI-07": {"user": "AI 에이전트", "english": "AI Agents", "tech": "agents", "status": "권장안"},
    "UI-08": {"user": "협업·의사결정", "english": "Collaboration and Decision", "tech": "collaboration", "status": "권장안"},
    "UI-09": {"user": "시스템 관리", "english": "System Administration", "tech": "administration", "status": "권장안"},
    "UI-10": {"user": "기준정보 관리", "english": "Master Data Management", "tech": "master_data_panel", "status": "권장안"},
    "UI-11": {"user": "코드·항목 연결표", "english": "Crosswalk", "tech": "crosswalk_panel", "status": "권장안"},
    "UI-16": {"user": "제작 단계 지도", "english": "Build Stage Map", "tech": "workflow_strip", "status": "과거 별칭"},
    "UI-19": {"user": "현재 작업 영역", "english": "Active Work Canvas", "tech": "adaptive_phase_canvas", "status": "권장안"},
    "UI-21": {"user": "결과 미리보기", "english": "Result Preview", "tech": "preview_panel", "status": "권장안"},
    "UI-24": {"user": "사용자 검토 요청", "english": "User Review Request", "tech": "user_review_request", "status": "권장안"},
    "UI-25": {"user": "Jarvis 패널", "english": "Jarvis Assistant Panel", "tech": "jarvis_rail", "status": "권장안"},
}

P0_RESOLUTIONS = {
    1: "Jarvis=제품 비서, 품질 감독 에이전트=파이프라인 판정자, Supervisor Daemon=기술 감시기로 분리",
    2: "메뉴는 ‘업무 앱 만들기’, 화면은 ‘업무 앱 스튜디오’, 기술 구조는 ‘Host-Governed Operational App Factory’",
    3: "MDM=관리체계, Master Data=관리대상, 기준정보=사용자 용어, Reference Data=코드·참조정보",
    4: "데이터 카탈로그=발견, 지식 허브=근거 활용, Reference Registry=승인 자료, Source Registry=원천 관리",
    5: "Tenant=고객 격리, Scope=권한 범위, Context=현재 선택, Workspace=협업 작업공간",
    6: "Program/Project는 사용자 업무 객체, Task/Sprint/Stage/Phase는 내부 실행 객체로 분리",
    7: "사용자 화면은 ‘사용자 검토 관문’, HOTL은 기술 문서와 코드에서만 사용",
    8: "Release=버전 확정, Delivery=사용자 전달, Promotion=조직 승격, Publication=보고서 발간, Export=파일 반출",
    9: "Digital Twin=제품 영역, Simulator=도구, Simulation=실행, Planning=계획 업무로 분리",
    10: "Baseline=비교 기준선, Snapshot=시점 데이터, Version=개정본, Checkpoint=실행 재개 지점",
    11: "ACTUAL은 업무 종류, SYNTHETIC은 출처. 두 축을 동시에 표시해 합성값의 실제값 오인을 차단",
    12: "App Manifest=앱 전체 선언, Capability Manifest=행동 권한, Data Contract=데이터 규약, Scope Contract=범위 규약",
    13: "Agent/Calculation/Traceability/Knowledge Graph를 목적별 고유명으로 사용하고 ‘Graph’ 단독 사용 금지",
    14: "Governance=정책·통제, Administration=운영 설정, Program Admin=앱 운영, Agent Governance=에이전트 자산 통제",
    15: "Enterprise Management Hub=첫 화면, Executive Cockpit=경영진 화면, Mega Boardroom=통합 프로젝트 관제",
}


def load_policy() -> dict[str, Any]:
    """사람이 검토하는 전환 정책을 읽어 코드 밖의 최종 권장안을 적용한다."""
    policy = json.loads(POLICY_SOURCE.read_text(encoding="utf-8"))
    overrides = policy.get("term_overrides", {})
    if not isinstance(overrides, dict):
        raise ValueError("term_overrides는 객체여야 합니다.")
    OVERRIDES.update(overrides)
    return policy


def _cells(line: str) -> list[str]:
    return [part.strip().replace("`", "") for part in line.strip().strip("|").split("|")]


def parse_inventory() -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    entries: list[dict[str, Any]] = []
    conflicts: list[dict[str, Any]] = []
    section = 0
    for line in SOURCE.read_text(encoding="utf-8").splitlines():
        heading = re.match(r"^##\s+(\d+)\.", line)
        if heading:
            section = int(heading.group(1))
            continue
        if not line.startswith("|") or re.match(r"^\|[\s:-]+\|", line):
            continue
        row = _cells(line)
        if section in AREA_BY_SECTION and len(row) >= 4 and re.fullmatch(r"[A-Z]+-\d+", row[0]):
            term_id, current, kind, issue = row[:4]
            override = OVERRIDES.get(term_id, {})
            if override:
                recommended = override["user"]
                english = override.get("english", current)
                technical = override.get("tech", current)
                status = override.get("status", "권장안")
            elif kind == "TECH":
                recommended, english, technical, status = current, current, current, "기술 전용"
            elif kind == "LEGACY":
                recommended, english, technical, status = "대체 용어 검토 중", current, current, "과거 별칭"
            elif kind in {"MIXED", "PLANNED"}:
                recommended, english, technical, status = current, current, current, "검토 필요"
            else:
                recommended, english, technical, status = current, current, current, "현행 유지"

            if override.get("guidance"):
                guidance = override["guidance"]
            elif status == "기술 전용":
                guidance = "사용자 화면에는 쉬운 업무 표현을 우선하고, 기술 설명·API·코드에서 이 명칭을 병기한다."
            elif status == "과거 별칭":
                guidance = "과거 문서 검색과 인수인계를 위해 보존하되 신규 사용자 화면에는 사용하지 않는다."
            elif status == "검토 필요":
                guidance = "현재 용어를 임시 유지한다. 의미 경계가 확정되기 전 일괄 치환하지 않는다."
            else:
                guidance = "사용자 화면은 권장 사용자 용어를 우선하고, 기술 설명에는 현재 용어와 기술 표준명을 병기한다."

            entries.append({
                "id": term_id,
                "area": AREA_BY_SECTION[section],
                "current_term": current,
                "current_kind": kind,
                "recommended_user_term": recommended,
                "recommended_english_term": english,
                "technical_canonical_name": technical,
                "definition": override.get("definition") or (
                    f"'{current}'은(는) {AREA_BY_SECTION[section]} 영역에서 현재 사용하는 "
                    f"{kind} 용어이며, 사용자 권장 표현은 '{recommended}'입니다."
                ),
                "redefinition_issue": issue,
                "migration_status": status,
                "usage_guidance": guidance,
                "legacy_alias": current if status in {"권장안", "과거 별칭"} and current != recommended else "",
                "source": str(SOURCE.relative_to(ROOT)).replace("\\", "/"),
            })
        elif section == 15 and len(row) >= 3 and row[0].isdigit():
            priority = int(row[0])
            conflicts.append({
                "priority": priority,
                "conflict_group": row[1],
                "decision_required": row[2],
                "recommended_resolution": P0_RESOLUTIONS.get(priority, "검토 필요"),
                "status": "팀 권장안",
            })
    return entries, conflicts


def build_json(
    entries: list[dict[str, Any]],
    conflicts: list[dict[str, Any]],
    policy: dict[str, Any],
) -> dict[str, Any]:
    policy_meta = policy["metadata"]
    return {
        "metadata": {
            "title": "AI Factory Studio 기술·제품 전환 용어 사전",
            "version": policy_meta["version"],
            "generated_on": policy_meta["generated_on"],
            "purpose": "현재 용어를 보존하면서 사용자 권장 용어와 기술 표준명을 함께 관리",
            "canonical_source": str(SOURCE.relative_to(ROOT)).replace("\\", "/"),
            "policy_source": str(POLICY_SOURCE.relative_to(ROOT)).replace("\\", "/"),
            "decision_state": policy_meta["decision_state"],
            "decision_label": policy_meta["decision_label"],
            "final_approval_required": policy_meta["final_approval_required"],
            "final_approver": policy_meta["final_approver"],
            "entry_count": len(entries),
            "conflict_count": len(conflicts),
            "policy": policy["policy"],
            "status_definitions": policy["status_definitions"],
        },
        "entries": entries,
        "p0_conflicts": conflicts,
    }


def write_json(payload: dict[str, Any]) -> None:
    text = json.dumps(payload, ensure_ascii=False, indent=2) + "\n"
    for path in (CANONICAL_JSON, FRONTEND_JSON):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")


def style_sheet(ws, widths: list[int]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx) if idx <= 26 else "A"].width = width
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="112E51")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    thin = Side(style="thin", color="D7DEE7")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=9, color="17243A")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
        ws.row_dimensions[row[0].row].height = 42


def add_table(ws, name: str) -> None:
    table = Table(displayName=name, ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)


def write_xlsx(payload: dict[str, Any]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "용어사전"
    headers = [
        "ID", "영역", "현재 용어", "현재 성격", "권장 화면 표기", "권장 영문명",
        "기술 표준명", "정의", "재정의 쟁점", "전환 상태", "사용 지침", "과거 별칭", "출처",
    ]
    ws.append(headers)
    for e in payload["entries"]:
        ws.append([
            e["id"], e["area"], e["current_term"], e["current_kind"],
            e["recommended_user_term"], e["recommended_english_term"],
            e["technical_canonical_name"], e["definition"], e["redefinition_issue"],
            e["migration_status"], e["usage_guidance"], e["legacy_alias"], e["source"],
        ])
    style_sheet(ws, [12, 27, 27, 12, 27, 29, 28, 34, 48, 14, 48, 27, 54])
    add_table(ws, "TerminologyDictionary")

    ws2 = wb.create_sheet("P0_충돌")
    ws2.append(["우선순위", "충돌 묶음", "결정해야 할 것", "초기 권장안", "상태"])
    for c in payload["p0_conflicts"]:
        ws2.append([c["priority"], c["conflict_group"], c["decision_required"], c["recommended_resolution"], c["status"]])
    style_sheet(ws2, [10, 48, 60, 74, 14])
    add_table(ws2, "P0TerminologyConflicts")

    ws3 = wb.create_sheet("분류기준")
    ws3.append(["분류", "의미", "기본 처리"])
    classifications = [
        ("PRODUCT", "고객·사용자·경영진에게 노출할 제품 언어", "권장 사용자 용어를 우선"),
        ("DOMAIN", "제조·경영·데이터의 업무 의미", "업무 정의와 영문 표준명을 함께 유지"),
        ("TECH", "API·코드·운영자 문서의 구현 언어", "사용자 UI 직접 노출을 피하고 기술 설명에 병기"),
        ("MIXED", "제품 언어와 구현 언어가 섞인 상태", "의미 경계 확정 전 일괄 치환 금지"),
        ("PLANNED", "기획은 있으나 구현·노출 수준 미확정", "검토 필요 상태로 관리"),
        ("LEGACY", "과거 문서·시안의 용어", "검색 별칭으로 보존하고 신규 UI에서 제외"),
    ]
    for row in classifications:
        ws3.append(row)
    style_sheet(ws3, [16, 56, 62])
    add_table(ws3, "TerminologyClassifications")

    ws4 = wb.create_sheet("사용지침")
    ws4.append(["구분", "내용"])
    guidance = [
        ("목적", "현재 기술 용어를 지우지 않고, 실제 사용자에게 보여 줄 권장 용어와 함께 관리하는 전환 사전입니다."),
        ("사전 버전", payload["metadata"]["version"]),
        ("결정 상태", f'{payload["metadata"]["decision_label"]} · 최종 승인자: {payload["metadata"]["final_approver"]}'),
        ("사용자 화면", "권장 화면 표기를 제목·메뉴·버튼에 우선 사용합니다. 쉬운 표현이 정확성을 떨어뜨리면 표준 용어를 유지합니다."),
        ("표준 약어", "WBS·RFP·PRD처럼 널리 통용되는 약어는 유지합니다. 최초 노출에는 WBS(작업분해구조)처럼 뜻을 병기하고 반복 노출에는 약어만 씁니다."),
        ("RFP 의미 구분", "현재 내부 산출물은 RFP(요구사항 정의서)로 설명합니다. 외부 공급자에게 제안을 요청하는 문서일 때만 RFP(제안요청서)라고 씁니다."),
        ("행동 문구", "표준 약어 자체를 없애지 않고 'WBS 작성 시작', 'PRD 검토', 'RFP 승인'처럼 사용자의 행동을 함께 표시합니다."),
        ("기술 설명", "현재 용어와 기술 표준명을 병기할 수 있습니다. 예: 사용자 검토 관문(User Review Gate, 내부 코드 HOTL)."),
        ("검토 필요", "확정되지 않은 항목입니다. 일괄 검색·치환하거나 API/DB 이름을 변경하지 않습니다."),
        ("과거 별칭", "기존 문서 검색과 인수인계를 위해 남깁니다. 신규 UI 문구로 다시 사용하지 않습니다."),
        ("정본", "Markdown 인벤토리를 고친 뒤 이 생성기를 실행합니다. Excel과 UI JSON을 직접 따로 수정하지 않습니다."),
        ("정책 파일", payload["metadata"]["policy_source"]),
        ("생성 명령", "venv\\Scripts\\python.exe scripts\\build_technology_terminology_dictionary.py"),
    ]
    for row in guidance:
        ws4.append(row)
    style_sheet(ws4, [18, 118])
    add_table(ws4, "TerminologyUsageGuide")

    ws5 = wb.create_sheet("변경대상")
    ws5.append(["ID", "현재 용어", "권장 화면 표기", "기술 표준명", "전환 상태"])
    for e in payload["entries"]:
        if e["migration_status"] not in {"현행 유지", "표준 약어 유지"}:
            ws5.append([e["id"], e["current_term"], e["recommended_user_term"], e["technical_canonical_name"], e["migration_status"]])
    style_sheet(ws5, [12, 34, 34, 34, 16])
    add_table(ws5, "TerminologyMigrationTargets")

    ws6 = wb.create_sheet("표준약어")
    ws6.append(["용어", "최초 노출", "반복 노출", "행동 문구 예시", "의미·주의사항"])
    standard_terms = [
        ("WBS", "WBS(작업분해구조)", "WBS", "WBS 작성 시작 · WBS 검토 · WBS 승인",
         "프로젝트 목표와 산출물을 실행 가능한 작업 단위로 계층적으로 분해한 구조"),
        ("RFP", "RFP(요구사항 정의서)", "RFP", "RFP 작성 · RFP 검토 · RFP 승인",
         "현재 시스템에서는 후속 기획·설계·검증의 기준 문서. 외부 발주 문서일 때만 RFP(제안요청서)"),
        ("PRD", "PRD(제품 요구사항 정의서)", "PRD", "PRD 작성 · PRD 검토 · PRD 승인",
         "RFP에서 확정한 요구를 제품 기능·사용 흐름·수용 기준으로 구체화한 문서"),
    ]
    for row in standard_terms:
        ws6.append(row)
    style_sheet(ws6, [16, 30, 20, 42, 88])
    add_table(ws6, "StandardAcronymGuide")

    for sheet in wb.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.outlinePr.summaryBelow = True
    XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX)


def validate(payload: dict[str, Any]) -> None:
    assert len(payload["entries"]) == 347, f"용어 수 불일치: {len(payload['entries'])}"
    assert len(payload["p0_conflicts"]) == 15, f"P0 충돌 수 불일치: {len(payload['p0_conflicts'])}"
    ids = [e["id"] for e in payload["entries"]]
    assert len(ids) == len(set(ids)), "중복 ID가 있습니다."
    assert all(e["current_term"] and e["recommended_user_term"] for e in payload["entries"])
    assert not any(e["migration_status"] == "검토 필요" for e in payload["entries"]), "검토 필요 용어가 남아 있습니다."
    assert not any(e["recommended_user_term"] == "대체 용어 검토 중" for e in payload["entries"]), "미정 대체어가 남아 있습니다."
    wb = load_workbook(XLSX, read_only=False, data_only=False)
    assert wb.sheetnames == ["용어사전", "P0_충돌", "분류기준", "사용지침", "변경대상", "표준약어"]
    assert wb["용어사전"].max_row == 348
    assert wb["P0_충돌"].max_row == 16
    assert wb["표준약어"].max_row == 4
    wb.close()


def main() -> None:
    policy = load_policy()
    entries, conflicts = parse_inventory()
    payload = build_json(entries, conflicts, policy)
    write_json(payload)
    write_xlsx(payload)
    validate(payload)
    print(f"OK: {len(entries)}개 용어, {len(conflicts)}개 P0 충돌")
    print(XLSX)


if __name__ == "__main__":
    main()
