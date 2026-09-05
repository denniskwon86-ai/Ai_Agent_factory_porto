#!/usr/bin/env python3
"""검증된 Starter Kit 계약에서 사용자 교체용 Excel 35종을 생성한다."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
KIT_ROOT = ROOT / "starter_kits" / "KIT-MFG-NONFERROUS-PROCUREMENT" / "1.0.0"
OUTPUT_ROOT = KIT_ROOT / "templates" / "excel"

NAVY = "1B2A41"
BLUE = "2F6FED"
RED = "C83232"
PALE_BLUE = "EAF2FF"
PALE_YELLOW = "FFF4CC"
PALE_GREEN = "E9F6EC"
LIGHT = "F3F5F7"
WHITE = "FFFFFF"
GRAY = "667085"
THIN = Side(style="thin", color="D9DEE7")


def read_csv(path: Path) -> List[Dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def style_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="111827")
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def title(ws, text: str, end_col: int = 6) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    c = ws.cell(1, 1, text)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.font = Font(name="Arial", size=18, bold=True, color=WHITE)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34


def header_row(ws, row: int, columns: int) -> None:
    for cell in ws[row][:columns]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)


def add_instructions(wb: Workbook, dataset: Dict[str, Any], contract: Dict[str, Any]) -> None:
    ws = wb.active
    ws.title = "INSTRUCTIONS"
    title(ws, f"{dataset['dataset_id']} · {dataset['name']} 입력 템플릿")
    rows = [
        ("키트", "KIT-MFG-NONFERROUS-PROCUREMENT 1.0.0"),
        ("용도", "검증된 가상기업 샘플을 참고하여 회사 데이터로 교체·등록"),
        ("중요", "이 파일의 예시값은 모두 SYNTHETIC이며 실제 경영 의사결정에 사용할 수 없습니다."),
        ("편집 순서", "1) DATA의 파란 예시행 확인 → 2) 노란 빈 행부터 입력 → 3) VALIDATION 확인 → 4) CHECKS가 PASS인지 확인"),
        ("필수 범위", "tenant_id·scope_node_id"
         + ("·cost_center_name" if any(
             field.get("name") == "cost_center_name" and field.get("required")
             for field in contract.get("schema", {}).get("fields", [])) else "")
         + "이 없으면 Fail-closed로 등록이 거부됩니다."),
        ("업무 키", ", ".join(contract.get("business_keys", []))),
        ("선행 데이터", ", ".join(contract.get("dependencies", [])) or "없음"),
        ("분류 원칙", "data_class=SYNTHETIC은 샘플의 출처 분류이며, business_data_kind는 값의 업무 의미(ACTUAL/PLAN 등)입니다."),
        ("사용자 입력", "노란색 셀은 회사 맞춤 입력 영역입니다. 예시행은 삭제하거나 별도 보관한 뒤 실제값으로 대체하십시오."),
    ]
    for idx, (label, value) in enumerate(rows, 3):
        ws.cell(idx, 1, label)
        ws.cell(idx, 2, value)
        ws.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=6)
        ws.cell(idx, 1).font = Font(name="Arial", size=10, bold=True, color=NAVY)
        ws.cell(idx, 1).fill = PatternFill("solid", fgColor=LIGHT)
        if label == "중요":
            ws.cell(idx, 2).font = Font(name="Arial", size=10, bold=True, color=RED)
    ws.column_dimensions["A"].width = 18
    for col in "BCDEF":
        ws.column_dimensions[col].width = 18
    ws.freeze_panes = "A3"
    style_sheet(ws)


def add_data(wb: Workbook, dataset_id: str, sample_rows: List[Dict[str, str]]) -> None:
    ws = wb.create_sheet("DATA")
    headers = list(sample_rows[0].keys())
    ws.append(headers)
    for row in sample_rows[:3]:
        ws.append([row.get(h, "") for h in headers])
    ws.append(["" for _ in headers])
    header_row(ws, 1, len(headers))
    for row in ws.iter_rows(min_row=2, max_row=4, max_col=len(headers)):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
            cell.font = Font(name="Arial", size=10, color=BLUE)
    for cell in ws[5][:len(headers)]:
        cell.fill = PatternFill("solid", fgColor=PALE_YELLOW)
        cell.font = Font(name="Arial", size=10, color=BLUE)
    ws["A5"].comment = Comment("이 행부터 회사 데이터를 입력하십시오. 예시행은 SYNTHETIC입니다.", "AI Factory Studio")
    table = Table(displayName=f"T_{dataset_id.replace('-', '_')}", ref=f"A1:{ws.cell(5, len(headers)).coordinate}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(5, len(headers)).coordinate}"
    for idx, name in enumerate(headers, 1):
        values = [len(str(name))] + [len(str(r.get(name, ""))) for r in sample_rows[:3]]
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = min(42, max(12, max(values) + 2))
    list_rules = {
        "data_class": '"ACTUAL,PLAN,FORECAST,SCENARIO,REFERENCE,SYNTHETIC"',
        "business_data_kind": '"ACTUAL,PLAN,FORECAST,SCENARIO,REFERENCE"',
        "quality_status": '"PASS,WARNING,QUARANTINE,REJECTED"',
        "certification_status": '"RAW,VALIDATED,CERTIFIED,CERTIFIED_FOR_DEMO"',
    }
    for field, formula in list_rules.items():
        if field in headers:
            col = ws.cell(1, headers.index(field) + 1).column_letter
            dv = DataValidation(type="list", formula1=formula, allow_blank=False)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}10000")
    style_sheet(ws)


def add_dictionary(wb: Workbook, contract: Dict[str, Any]) -> None:
    ws = wb.create_sheet("DATA_DICTIONARY")
    ws.append(["필드명", "형식", "필수", "업무 키", "설명"])
    for field in contract["schema"]["fields"]:
        ws.append([field["name"], field["type"], "Y" if field["required"] else "N",
                   "Y" if field["business_key"] else "N", field["description"]])
    header_row(ws, 1, 5)
    widths = [28, 14, 10, 12, 54]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = width
    ws.freeze_panes = "A2"
    style_sheet(ws)


def add_code_map(wb: Workbook) -> None:
    ws = wb.create_sheet("CODE_MAP")
    ws.append(["필드", "허용값", "의미"])
    values = [
        ("data_class", "SYNTHETIC", "가상·합성 데이터. Actual 승격 금지"),
        ("business_data_kind", "ACTUAL", "업무상 실적 형태의 값(샘플에서도 출처는 SYNTHETIC)"),
        ("business_data_kind", "PLAN", "확정 또는 작업 중인 계획"),
        ("business_data_kind", "FORECAST", "모델·사용자 가정 기반 전망"),
        ("business_data_kind", "SCENARIO", "What-if 대안"),
        ("business_data_kind", "REFERENCE", "마스터·기준·공식 외부지표"),
        ("quality_status", "PASS", "계약·참조·대사 통과"),
        ("quality_status", "QUARANTINE", "운영 영역에 넣지 않고 사유와 함께 격리"),
        ("certification_status", "CERTIFIED_FOR_DEMO", "데모 용도 검증 완료. 실제 의사결정용 아님"),
    ]
    for row in values:
        ws.append(row)
    header_row(ws, 1, 3)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 62
    style_sheet(ws)


def add_validation(wb: Workbook, contract: Dict[str, Any]) -> None:
    ws = wb.create_sheet("VALIDATION")
    ws.append(["검증 ID", "검증 내용", "오류 시 처리", "근거"])
    rules = [
        ("VAL-01", "필수 공통 필드가 모두 존재", "REJECT", "데이터 계약"),
        ("VAL-02", "tenant_id·scope_node_id가 비어 있지 않음", "REJECT", "Fail-closed"),
        ("VAL-03", "업무 키가 중복되지 않음", "QUARANTINE", ", ".join(contract.get("business_keys", []))),
        ("VAL-04", "선행 데이터의 참조 키가 존재", "QUARANTINE", ", ".join(contract.get("dependencies", [])) or "없음"),
        ("VAL-05", "SYNTHETIC 표기가 Actual 출처로 변경되지 않음", "REJECT", "분류 정책"),
        ("VAL-06", "기간·단위·통화가 코드맵과 일치", "QUARANTINE", "MDM·Crosswalk"),
    ]
    if any(field.get("name") == "cost_center_name" and field.get("required")
           for field in contract.get("schema", {}).get("fields", [])):
        rules.append(("VAL-07", "같은 cost_center_id의 명칭·tenant·scope가 모두 일치",
                      "REJECT", "원가센터 집합 정본"))
    for row in rules:
        ws.append(row)
    header_row(ws, 1, 4)
    for col, width in zip("ABCD", [14, 54, 18, 40]):
        ws.column_dimensions[col].width = width
    style_sheet(ws)


def add_checks(wb: Workbook, contract: Dict[str, Any]) -> None:
    ws = wb.create_sheet("CHECKS")
    ws.append(["지표", "결과", "판정 기준"])
    ws.append(["입력 행 수", "=COUNTA(DATA!A:A)-1", "1건 이상이면 입력 데이터 존재"])
    ws.append(["SYNTHETIC 표시 행 수", '=COUNTIF(DATA!D:D,"SYNTHETIC")', "샘플 단계에서는 입력 행 수와 동일"])
    ws.append(["분류 일치", '=IF(B2=B3,"PASS","REVIEW")', "PASS여야 함"])
    fields = [str(field.get("name") or "")
              for field in contract.get("schema", {}).get("fields", [])]
    if "cost_center_name" in fields:
        name_col = get_column_letter(fields.index("cost_center_name") + 1)
        missing_formula = (f'=COUNTBLANK(DATA!C2:INDEX(DATA!C:C,B2+1))+'
                           f'COUNTBLANK(DATA!{name_col}2:INDEX(DATA!{name_col}:{name_col},B2+1))')
        ws.append(["필수 범위·원가센터 명칭 누락", missing_formula, "0이어야 함"])
    else:
        ws.append(["범위 누락", '=COUNTBLANK(DATA!C2:INDEX(DATA!C:C,B2+1))', "0이어야 함"])
    ws.append(["종합", '=IF(AND(B4="PASS",B5=0),"PASS","REVIEW")', "PASS 후 등록"])
    header_row(ws, 1, 3)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 50
    for cell in (ws["B2"], ws["B3"], ws["B4"], ws["B5"], ws["B6"]):
        cell.fill = PatternFill("solid", fgColor=PALE_GREEN)
    ws.conditional_formatting.add("B4:B6", CellIsRule(operator="equal", formula=['"REVIEW"'],
                                                        fill=PatternFill("solid", fgColor="FDECEC")))
    style_sheet(ws)


def build_index(manifest: Dict[str, Any]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "SUMMARY"
    title(ws, "AI Factory Studio · Sample Company Starter Kit", 8)
    info = [
        ("키트", manifest["kit_id"]), ("버전", manifest["version"]),
        ("회사", manifest["company_name"] + " (가상)"), ("상태", manifest["status"]),
        ("데이터셋", manifest["dataset_count"]), ("주의", "SYNTHETIC · 실제 경영 의사결정 금지"),
    ]
    for idx, (label, value) in enumerate(info, 3):
        ws.cell(idx, 1, label)
        ws.cell(idx, 2, value)
        ws.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=8)
        ws.cell(idx, 1).fill = PatternFill("solid", fgColor=LIGHT)
        ws.cell(idx, 1).font = Font(name="Arial", bold=True, color=NAVY)
    ws.column_dimensions["A"].width = 18
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 16
    cat = wb.create_sheet("DATA_CATALOG")
    cat.append(["ID", "데이터셋", "필수", "업무 키", "선행 데이터", "Excel 파일"])
    for dataset in manifest["datasets"]:
        cat.append([dataset["dataset_id"], dataset["name"], "Y" if dataset["required"] else "N",
                    ", ".join(dataset["keys"]), ", ".join(dataset["deps"]), f"{dataset['dataset_id']}.xlsx"])
    header_row(cat, 1, 6)
    for col, width in zip("ABCDEF", [14, 34, 10, 34, 34, 20]):
        cat.column_dimensions[col].width = width
    cat.freeze_panes = "A2"
    profiles_ws = wb.create_sheet("COMPANY_PROFILES")
    profiles_ws.append(["Profile ID", "회사명", "유형", "업종", "목적", "유효기간", "추천 앱", "기본 가정"])
    for profile_path in sorted((KIT_ROOT / "company_profiles").glob("*.json")):
        profile = json.loads(profile_path.read_text(encoding="utf-8"))
        profiles_ws.append([
            profile["company_profile_id"], profile["company_name"], profile["profile_role"],
            f"{profile['industry_code']} · {profile['industry_name']}", profile["purpose"],
            profile["valid_until"], ", ".join(profile["recommended_apps"]),
            json.dumps(profile["default_assumptions"], ensure_ascii=False, sort_keys=True),
        ])
    header_row(profiles_ws, 1, 8)
    for col, width in zip("ABCDEFGH", [38, 34, 28, 34, 60, 16, 24, 60]):
        profiles_ws.column_dimensions[col].width = width
    profiles_ws.freeze_panes = "A2"
    style_sheet(ws)
    style_sheet(cat)
    style_sheet(profiles_ws)
    path = OUTPUT_ROOT / "AFS_Starter_Kit_Onboarding_Index.xlsx"
    wb.save(path)
    return path


def main() -> None:
    manifest_path = KIT_ROOT / "manifest.json"
    if not manifest_path.exists():
        raise SystemExit("검증된 Starter Kit manifest가 없습니다.")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest.get("status") != "VALIDATED_FOR_DEMO":
        raise SystemExit(f"검증 완료 상태가 아닙니다: {manifest.get('status')}")
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    generated = []
    for dataset in manifest["datasets"]:
        dataset_id = dataset["dataset_id"]
        contract = json.loads((KIT_ROOT / "contracts" / f"{dataset_id}.contract.json").read_text(encoding="utf-8"))
        samples = read_csv(KIT_ROOT / "samples" / "quick" / f"{dataset_id}.csv")
        wb = Workbook()
        add_instructions(wb, dataset, contract)
        add_data(wb, dataset_id, samples)
        add_dictionary(wb, contract)
        add_code_map(wb)
        add_validation(wb, contract)
        add_checks(wb, contract)
        path = OUTPUT_ROOT / f"{dataset_id}.xlsx"
        wb.save(path)
        generated.append(path)
    index_path = build_index(manifest)
    print(json.dumps({"status": "generated", "dataset_workbooks": len(generated),
                      "index": str(index_path), "output_root": str(OUTPUT_ROOT)},
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
