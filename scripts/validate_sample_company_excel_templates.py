#!/usr/bin/env python3
"""Starter Kit Excel 35종의 구조·수식·가독성 계약을 검증한다."""

from __future__ import annotations

import json
import hashlib
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
KIT_ROOT = ROOT / "starter_kits" / "KIT-MFG-NONFERROUS-PROCUREMENT" / "1.0.0"
EXCEL_ROOT = KIT_ROOT / "templates" / "excel"
REQUIRED_SHEETS = ["INSTRUCTIONS", "DATA", "DATA_DICTIONARY", "CODE_MAP", "VALIDATION", "CHECKS"]


def main() -> None:
    manifest = json.loads((KIT_ROOT / "manifest.json").read_text(encoding="utf-8"))
    failures = []
    formula_count = 0
    for dataset in manifest["datasets"]:
        dataset_id = dataset["dataset_id"]
        path = EXCEL_ROOT / f"{dataset_id}.xlsx"
        if not path.exists():
            failures.append(f"{dataset_id}:파일없음")
            continue
        formula_wb = load_workbook(path, data_only=False, read_only=False)
        value_wb = load_workbook(path, data_only=True, read_only=False)
        if formula_wb.sheetnames != REQUIRED_SHEETS:
            failures.append(f"{dataset_id}:시트={formula_wb.sheetnames}")
        data_ws = formula_wb["DATA"]
        if data_ws.freeze_panes != "A2":
            failures.append(f"{dataset_id}:DATA고정창")
        if not data_ws.tables:
            failures.append(f"{dataset_id}:DATA표없음")
        if data_ws["A5"].fill.fgColor.rgb not in {"00FFF4CC", "FFFFF4CC", "FFF4CC"}:
            failures.append(f"{dataset_id}:입력행표시")
        for ws in formula_wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    # LibreOffice/Windows는 Arial에 없는 한글 글리프를 맑은 고딕으로
                    # 정상 대체한다. 두 글꼴 모두 기업 문서용 전문 글꼴로 허용한다.
                    if cell.value is not None and cell.font.name not in {"Arial", "맑은 고딕", None}:
                        failures.append(f"{dataset_id}:{ws.title}!{cell.coordinate}:font={cell.font.name}")
                        break
        check_formula_ws = formula_wb["CHECKS"]
        check_value_ws = value_wb["CHECKS"]
        for coord in ("B2", "B3", "B4", "B5", "B6"):
            if not isinstance(check_formula_ws[coord].value, str) or not check_formula_ws[coord].value.startswith("="):
                failures.append(f"{dataset_id}:수식없음:{coord}")
            else:
                formula_count += 1
            cached = check_value_ws[coord].value
            if cached is None or (isinstance(cached, str) and cached.startswith("#")):
                failures.append(f"{dataset_id}:수식결과오류:{coord}={cached}")
        if check_value_ws["B6"].value != "PASS":
            failures.append(f"{dataset_id}:종합={check_value_ws['B6'].value}")
        formula_wb.close()
        value_wb.close()

    index_path = EXCEL_ROOT / "AFS_Starter_Kit_Onboarding_Index.xlsx"
    if not index_path.exists():
        failures.append("INDEX:파일없음")
    else:
        index_wb = load_workbook(index_path, read_only=True)
        if index_wb.sheetnames != ["SUMMARY", "DATA_CATALOG", "COMPANY_PROFILES"]:
            failures.append(f"INDEX:시트={index_wb.sheetnames}")
        if index_wb["DATA_CATALOG"].max_row != 36:
            failures.append(f"INDEX:카탈로그행={index_wb['DATA_CATALOG'].max_row}")
        if index_wb["COMPANY_PROFILES"].max_row != 7:
            failures.append(f"INDEX:회사프로필행={index_wb['COMPANY_PROFILES'].max_row}")
        index_wb.close()

    result = {
        "status": "PASS" if not failures else "FAIL",
        "dataset_workbooks": len(manifest["datasets"]),
        "formula_cells_checked": formula_count,
        "failures": failures,
    }
    report = KIT_ROOT / "validations" / "excel_validation_report.json"
    report.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    if not failures:
        def digest(path: Path) -> str:
            value = hashlib.sha256()
            with path.open("rb") as stream:
                for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                    value.update(chunk)
            return value.hexdigest()

        manifest["excel_templates"] = {
            "status": "PASS",
            "dataset_workbooks": 35,
            "index_path": "templates/excel/AFS_Starter_Kit_Onboarding_Index.xlsx",
            "validation_report": report.relative_to(KIT_ROOT).as_posix(),
            "formula_cells_checked": formula_count,
        }
        manifest["file_index"] = [
            item for item in manifest.get("file_index", [])
            if item.get("profile") not in {"excel-template", "excel-index"}
        ]
        for dataset in manifest["datasets"]:
            path = EXCEL_ROOT / f"{dataset['dataset_id']}.xlsx"
            manifest["file_index"].append({
                "path": path.relative_to(KIT_ROOT).as_posix(),
                "sha256": digest(path),
                "dataset_id": dataset["dataset_id"],
                "profile": "excel-template",
            })
        manifest["file_index"].append({
            "path": index_path.relative_to(KIT_ROOT).as_posix(),
            "sha256": digest(index_path),
            "dataset_id": "ONBOARDING-INDEX",
            "profile": "excel-index",
        })
        (KIT_ROOT / "manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
